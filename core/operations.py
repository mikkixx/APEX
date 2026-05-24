import sys
import re
import bcrypt
from peewee import fn, DoesNotExist, IntegrityError, OperationalError
from datetime import date, timedelta, datetime
from db.models import User, ReadinessStatus, TrainingDiary, Recommendation, MedicalExam, MedicalMetric, TrainingPlan, Session, Message, SpecialistBinding
from db.connection import db
import os
import unicodedata

import sys as _sys

def _get_reports_dir():
    if hasattr(_sys, '_MEIPASS'):
        base = _sys.executable.replace(_sys.executable.split('\\')[-1], '')
    else:
        import os as _os
        base = _os.path.abspath('.')
    import os as _os
    path = _os.path.join(base, 'reports')
    _os.makedirs(path, exist_ok=True)
    return path


def _is_valid_email(email: str) -> bool:
    return bool(re.match(r"^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$", email))

def register(last_name, first_name, middle_name, email, password, confirm_password, role, specialization):
    if not all ([last_name, first_name, email, password, confirm_password, role, specialization]):
        return False, 'Заполните все обязательные поля', None
    if not _is_valid_email(email):
        return False, 'Неверный формат email', None
    if len(password) < 6:
        return False, 'Пароль должен содержать не менее 6 символов', None
    if password != confirm_password:
        return False, 'Пароли не совпадают', None
    
    try:
        if User.select().where(User.email == email).exists():
            return False, 'Пользователь с таким email уже существует. Пожалуйста, войдите в систему.', None
        
        password_hash = bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        User.create(
            last_name=last_name,
            first_name=first_name,
            middle_name=middle_name or None,
            email=email,
            password_hash=password_hash,
            role=role,
            specialization=specialization
        )
        return True, 'Регистрация прошла успешно. Перейдите к авторизации.', None
    except OperationalError as e:
        return False, f"Ошибка подключения к БД: {e}", None
    
def login(email, password):
    if not email or not password:
        return False, 'Заполните поля email и пароль', None
    
    try:
        user = User.get(User.email == email)

        if bcrypt.checkpw(password.encode('utf-8'), user.password_hash.encode('utf-8')):
            user_data = {
                'id': user.id,
                'last_name': user.last_name,
                'first_name': user.first_name,
                'middle_name': user.middle_name or "",
                'role': user.role,
                'specialization': user.specialization
            }
            return True, 'Вход выполнен успешно', user_data
        else:
            return False, 'Неверный пароль', None
        
    except DoesNotExist:
        return False, 'Пользователь с таким email не найден. Пожалуйста, зарегистрируйтесь.', None
    except OperationalError as e:
        return False, f"Ошибка подключения: {e}", None
def change_password(user_id, current_password, new_password, confirm_password):
    if not all([current_password, new_password, confirm_password]):
        return False, 'Заполните все поля', None
    if not (6 <= len(new_password) <= 32):
        return False, 'Пароль должен быть от 6 до 32 символов', None
    if not (re.search(r'[A-Z]', new_password) and re.search(r'[a-z]', new_password) and re.search(r'\d', new_password)):
        return False, 'Пароль должен содержать буквы разного регистра и цифры', None
    if new_password != confirm_password:
        return False, 'Пароли не совпадают', None
    if new_password == current_password:
        return False, 'Новый пароль не должен совпадать с текущим', None
    
    try:
        user = User.get_by_id(user_id)
        if not bcrypt.checkpw(current_password.encode('utf-8'), user.password_hash.encode('utf-8')):
            return False, 'Неверный текущий пароль', None
        
        new_hash = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
        user.password_hash = new_hash
        user.save()
        return True, 'Пароль успешно изменен', None
    except DoesNotExist:
        return False, 'Пользователь не найден', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    
def get_profile(user_id):
    try:
        if db.is_closed():
            db.connect()

        user = User.get_by_id(user_id)
        profile_data = {
            'id': user.id,
            'last_name': user.last_name,
            'first_name': user.first_name,
            'middle_name': user.middle_name or "",
            'role': user.role,
            'specialization': user.specialization,
            'photo_path': user.photo_path
        }

        if user.role == 'спортсмен':
            try:
                latest = ReadinessStatus.select().where(ReadinessStatus.athlete_id == user_id).order_by(ReadinessStatus.id.desc()).get()
                profile_data['current_status'] = latest.current_status
            except DoesNotExist:
                profile_data['current_status'] = 'Не установлен'

        return True, 'Данные профиля загружены', profile_data
    except DoesNotExist:
        return False, 'Профиль не найден', None
    except OperationalError as e:
        return False, f"Ошибка подключения к БД: {e}", None
    except Exception as e:
        return False, f"Ошибка в get_profile: {e}", None
                                                   
_PHOTO_UNSET = object()

def edit_profile(user_id, last_name, first_name, middle_name, email, specialization, photo_path=_PHOTO_UNSET):
    if not all([last_name, first_name, email, specialization]):
        return False, 'Заполните все обязательные поля', None
    if not _is_valid_email(email):
        return False, 'Неверный формат email', None
    
    try:
        if db.is_closed():
            db.connect()
        
        user = User.get_by_id(user_id)

        if User.select().where((User.email == email) & (User.id != user_id)).exists():
            return False, 'Этот email уже занят другим пользователем', None
        
        with db.atomic():
            user.last_name = last_name
            user.first_name = first_name
            user.middle_name = middle_name or None
            user.email = email
            user.specialization = specialization
            if photo_path is not _PHOTO_UNSET:
                user.photo_path = photo_path
            user.save()

        return True, 'Профиль успешно обновлен', None
    except DoesNotExist:
        return False, 'Пользователь не найден', None
    except OperationalError as e:
        return False, f"Ошибка подключения к БД: {e}", None
    
def _get_current_week():
    today = date.today()
    monday = today - timedelta(days=today.weekday())
    sunday = monday + timedelta(days=6)
    return monday, sunday
def add_diary_entry(athlete_id, entry_date, activity_type, duration, steps, sleep_hours, fatigue, mood, comment=None):
    if not all([entry_date, activity_type, duration, sleep_hours, fatigue, mood]):
        return False, 'Заполните все обязательные поля', None
    if not (1 <= duration <= 300):
        return False, 'Длительность занятия должна быть от 1 до 300 мин', None
    if steps < 0:
        return False, 'Количество шагов не может быть отрицательным числом', None
    if not (1 <= fatigue <= 10):
        return False, 'Усталость: шкала от 1 до 10', None
    if not (1 <= mood <= 10):
        return False, 'Настроение: шкала от 1 до 10', None
    if entry_date > date.today():
        return False, 'Дата не может быть в будущем', None
    
    clean_comment = comment.strip() if comment else None
    
    try:
        if db.is_closed():
            db.connect()

        existing = TrainingDiary.select().where(
            (TrainingDiary.athlete == athlete_id) &
            (TrainingDiary.date == entry_date) &
            (TrainingDiary.is_deleted == False)
        ).first()

        if existing:
            return False, f"Запись за {entry_date} уже существует. Отредактируйте ее.", {"existing_id": existing.id}
        
        with db.atomic():
            TrainingDiary.create(
                athlete = athlete_id,
                date = entry_date,
                activity_type = activity_type,
                duration = duration,
                steps = steps,
                sleep_hours = sleep_hours,
                fatigue = fatigue,
                mood = mood,
                comment = clean_comment
            )
        return True, 'Запись добавлена', None
    except OperationalError as e:
        return False, f"Ошибка подключения к БД: {e}", None
    
def edit_diary_entry(entry_id, athlete_id, entry_date, activity_type, duration, steps, sleep_hours, fatigue, mood, comment=None):
    if not all([entry_date, activity_type, duration, sleep_hours, fatigue, mood]):
        return False, 'Заполните все обязательные поля', None
    if not (1 <= duration <= 300):
        return False, 'Длительность занятия должна быть от 1 до 300 мин', None
    if steps < 0:
        return False, 'Количество шагов не может быть отрицательным числом', None
    if not (1 <= fatigue <= 10):
        return False, 'Усталость: шкала от 1 до 10', None
    if not (1 <= mood <= 10):
        return False, 'Настроение: шкала от 1 до 10', None
    if entry_date > date.today():
        return False, 'Дата не может быть в будущем', None
    
    clean_comment = comment.strip() if comment else None

    try:
        entry = TrainingDiary.get_by_id(entry_id)
        if entry.athlete_id != athlete_id:
            return False, 'Доступ запрещен', None
        
        critical_changed = (entry.date != entry_date) or (entry.activity_type != activity_type)

        with db.atomic():
                entry.date = entry_date
                entry.activity_type = activity_type
                entry.duration = duration
                entry.steps = steps
                entry.sleep_hours = sleep_hours
                entry.fatigue = fatigue
                entry.mood = mood
                entry.comment = clean_comment
                entry.save()

                if critical_changed:
                    Recommendation.delete().where(
                        (Recommendation.linked_entity == 'дневник нагрузок') &
                        (Recommendation.linked_entity_id == entry_id)
                    ).execute()
            
        return True, 'Изменения сохранены', {'needs_review': critical_changed}
    except DoesNotExist:
        return False, 'Запись не найдена', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    
def delete_diary_entry(entry_id, athlete_id):
    try:
        entry = TrainingDiary.get_by_id(entry_id)
        if entry.athlete_id != athlete_id:
            return False, 'Доступ запрещен', None
        
        with db.atomic():
            entry.is_deleted = True
            entry.save()
        return True, 'Запись удалена', None
    except DoesNotExist:
        return False, 'Запись не найдена', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    
def get_medical_data(athlete_id, exam_date=None, exam_type=None):
    try:
        if db.is_closed():
            db.connect()

        conditions = [MedicalExam.athlete_id == athlete_id]

        if exam_date:
            if exam_date > date.today():
                return False, 'Дата не может быть в будущем', None
            conditions.append(MedicalExam.exam_date <= exam_date)
        
        if exam_type and exam_type.strip():
            conditions.append(MedicalExam.exam_type == exam_type.strip())

        exams = MedicalExam.select().where(*conditions).order_by(MedicalExam.exam_date.desc())
        result = []

        for exam in exams:
            metrics = MedicalMetric.select().where(MedicalMetric.exam_id == exam.id).order_by(MedicalMetric.id.desc())
            recommendations = Recommendation.select().where(
                (Recommendation.linked_entity == 'медкарта') & (Recommendation.linked_entity_id == exam.id)
            ).order_by(Recommendation.id.desc())

            try:
                doctor = exam.doctor
                doctor_fio = f"{doctor.last_name} {doctor.first_name} {doctor.middle_name or ''}".strip()
                doctor_email = doctor.email
            except DoesNotExist:
                doctor_fio = "Врач удалён"
                doctor_email = "—"

            result.append({ 
                'exam_date': exam.exam_date,
                'exam_type': exam.exam_type,
                'doctor_fio': doctor_fio,
                'doctor_email': doctor_email,
                'metrics': [
                    {'type': m.metric_type, 'value': m.value, 'unit': m.unit,
                     'ref_range': m.ref_range or '', 'is_critical': bool(m.is_critical)} 
                    for m in metrics
                ],
                'recommendations': list(recommendations),
                'exam': exam
            })

        return True, 'Медицинские данные загружены', result
    except OperationalError as e:
        return False, f"Ошибка подключения: {e}", None
    except Exception as e:
        return False, f"Неожиданная ошибка в get_medical_data: {e}", None
    
def get_training_plan(athlete_id, start_date=None, end_date=None):
    try:
        if db.is_closed():
            db.connect()

        if start_date is None or end_date is None:
            start_date, end_date = _get_current_week()

        plans = TrainingPlan.select().where(
            (TrainingPlan.athlete == athlete_id) &
            (TrainingPlan.start_date <= end_date) &
            (TrainingPlan.end_date >= start_date) &
            (TrainingPlan.is_deleted == False)
        )

        result = []
        for plan in plans:
            sessions = Session.select().where(
                (Session.plan == plan) &
                (Session.is_deleted == False) &
                (Session.date >= start_date) &
                (Session.date <= end_date)
            ).order_by(Session.date.asc(), Session.time.asc())

            result.append({'plan': plan, 'sessions': list(sessions)})
        return True, 'План загружен', result
    except OperationalError as e:
        return False, f"Ошибка подключения: {e}", None
    
def sync_overdue_sessions(athlete_id):
    try:
        if db.is_closed():
            db.connect()
        
        from datetime import datetime, timedelta
        cutoff = datetime.now() - timedelta(hours=24)
        
        overdue = Session.select().join(TrainingPlan).where(
            (TrainingPlan.athlete == athlete_id) &
            (Session.status == 'запланировано') &
            (
                (Session.date < cutoff.date()) |
                (
                    (Session.date == cutoff.date()) &
                    (Session.time is not None) &
                    (
                        datetime.combine(Session.date, Session.time) + timedelta(hours=24) < datetime.now()
                    )
                )
            ) &
            (Session.is_deleted == False)
        )
        
        updated = 0
        for s in overdue:
            s.status = 'пропущено'
            s.save()
            updated += 1
        
        return True, f'Обновлено {updated} занятий', None
    except OperationalError as e:
        return False, f'Ошибка БД: {e}', None
    except Exception as e:
        return False, f'Ошибка: {e}', None

def update_session_status(session_id, athlete_id, new_status):
    valid_statuses = ['запланировано', 'выполнено', 'пропущено']
    if new_status not in valid_statuses:
        return False, 'Недопустимый статус', None
    
    try:
        session = Session.get_by_id(session_id)
        if session.plan.athlete_id != athlete_id:
            return False, 'Доступ запрещен', None
        
        today = date.today()
        if session.date < today - timedelta(days=1) and new_status != 'пропущено':
            return False, 'Изменение статуса доступно только в день занятия или в течение 24 часов после', None
        
        with db.atomic():
            session.status = new_status
            session.save()
        return True, 'Статус сохранен', None
    except DoesNotExist:
        return False, 'Занятие не найдено', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    
def get_chat_partners(current_user_id):
    try:
        if db.is_closed():
            db.connect()
    
        sent_ids = [m.receiver_id for m in Message.select(Message.receiver_id).where(Message.sender_id == current_user_id).distinct()]
        recv_ids = [m.sender_id for m in Message.select(Message.sender_id).where(Message.receiver_id == current_user_id).distinct()]
        partner_ids = list(set(sent_ids + recv_ids))

        if not partner_ids:
            return True, 'Список пуст', []
        
        partners = User.select().where(User.id << partner_ids).order_by(User.last_name)
        result = [
            {'id': p.id, 'full_name': f"{p.last_name} {p.first_name} {p.middle_name or ''}".strip(), 'photo_path': p.photo_path, 'role': p.role, 'specialization': p.specialization} 
            for p in partners
        ]
        return True, 'Список загружен', result
    except OperationalError as e:
        return False, f"Ошибка подключения: {e}", None
    except Exception as e:
        return False, f"Ошибка загрузки собеседников: {e}", None

def get_chat_messages(current_user_id, partner_id, search_query=None, start_date=None, end_date=None):
    try:
        if db.is_closed():
            db.connect()
        
        conditions = [
            ((Message.sender == current_user_id) & (Message.receiver == partner_id)) |
            ((Message.sender == partner_id) & (Message.receiver == current_user_id))
        ]

        if start_date and end_date:
            conditions.append(Message.sent_at >= datetime.combine(start_date, datetime.min.time()))
            conditions.append(Message.sent_at <= datetime.combine(end_date, datetime.max.time()))

        if search_query and search_query.strip():
            clean_q = search_query.strip()
            if len(clean_q) < 3:
                return False, 'Поисковый запрос должен содержать минимум 3 символа', None
            conditions.append(fn.LOWER(Message.text).contains(clean_q.lower()))

        messages = (Message.select(Message, User)
            .join(User, on=(Message.sender == User.id))
            .where(*conditions)
            .order_by(Message.sent_at.asc()))
            
        result = []
        for msg in messages:
            is_mine = (msg.sender.id == current_user_id)
            result.append({
                'id': msg.id, 'text': msg.text, 'sent_at': msg.sent_at,
                'is_mine': is_mine, 'sender_name': 'Вы' if is_mine else f"{msg.sender.last_name} {msg.sender.first_name}"
            })
        return True, 'Сообщения загружены', result
    except OperationalError as e:
        return False, f"Ошибка подключения: {e}", None
    except Exception as e:
        return False, f"Ошибка загрузки сообщений: {e}", None

def send_message(sender_id, receiver_id, text):
    clean_text = text.strip()
    if not clean_text:
        return False, 'Сообщение не может быть пустым', None
    if len(clean_text) > 500:
        return False, 'Сообщение не может превышать 500 символов', None

    try:
        with db.atomic():
            Message.create(sender=sender_id, receiver=receiver_id, text=clean_text, sent_at=datetime.now()) # 🔹 Исправлено sent_at
        return True, 'Сообщение отправлено', None
    except OperationalError as e:
        return False, f"Ошибка отправки: {e}", None   

def add_chat_partner(current_user_id, target_email):
    if not target_email or not target_email.strip():
        return False, 'Введите email собеседника', None
    clean_email = target_email.strip()
    
    try:
        if db.is_closed():
            db.connect()
        try:
            target_user = User.get(User.email == clean_email)
        except DoesNotExist:
            return False, 'Пользователь с таким email не найден', None

        if target_user.id == current_user_id:
            return False, 'Нельзя добавить самого себя в собеседники', None
        
        return True, 'Собеседник добавлен', {
            'id': target_user.id, 'full_name': f"{target_user.last_name} {target_user.first_name}".strip(),
            'photo_path': target_user.photo_path, 'role': target_user.role, 'specialization': target_user.specialization
        }
    except OperationalError as e:
        return False, f"Ошибка подключения к базе: {e}", None  

def delete_account(user_id):
    try:
        user = User.get_by_id(user_id)
        if hasattr(User, 'is_active'):
            user.is_active = False
        else:
            user.email = f"disabled_{user.id}@deleted.local"
        user.save()
        db.close()
        return True, 'Аккаунт деактивирован. Сессия завершена.', None
    except DoesNotExist:
        return False, 'Пользователь не найден', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    

def get_my_athletes(specialist_id, page=1, per_page=5, search=None, sport_type=None, status=None):
    try:
        if db.is_closed():
            db.connect()

        binding_ids = list(
            SpecialistBinding.select(SpecialistBinding.athlete)
            .where(
                (SpecialistBinding.specialist == specialist_id) &
                (SpecialistBinding.status == 'активна') &
                (SpecialistBinding.is_deleted == False)
            ).tuples()
        )
        athlete_ids = [row[0] for row in binding_ids]

        if not athlete_ids:
            return True, 'Список пуст', {'athletes': [], 'total': 0, 'page': page, 'per_page': per_page}

        query = User.select().where(User.id << athlete_ids)

        if search and len(search.strip()) >= 3:
            q = search.strip().lower()
            query = query.where(
                (fn.LOWER(User.last_name).contains(q)) |
                (fn.LOWER(User.first_name).contains(q)) |
                (fn.LOWER(User.middle_name).contains(q))
            )

        if sport_type and sport_type.strip():
            query = query.where(fn.LOWER(User.specialization).contains(sport_type.strip().lower()))

        if status and status.strip():
            status_val = status.strip()
            valid_ids = []
            for aid in athlete_ids:
                try:
                    s = ReadinessStatus.select().where(ReadinessStatus.athlete == aid).order_by(ReadinessStatus.id.desc()).get()
                    if s.current_status == status_val:
                        valid_ids.append(aid)
                except DoesNotExist:
                    pass
            query = query.where(User.id << valid_ids)

        total = query.count()

        users = list(query.order_by(User.last_name.asc()).paginate(page, per_page))

        result = []
        for athlete in users:
            try:
                latest_status = ReadinessStatus.select().where(
                    ReadinessStatus.athlete == athlete.id
                ).order_by(ReadinessStatus.id.desc()).get()
                current_status = latest_status.current_status
            except DoesNotExist:
                current_status = 'Не установлен'

            result.append({
                'id': athlete.id,
                'last_name': athlete.last_name,
                'first_name': athlete.first_name,
                'middle_name': athlete.middle_name or '',
                'specialization': athlete.specialization,
                'current_status': current_status,
                'email': athlete.email,
                'photo_path': athlete.photo_path
            })

        return True, 'Список загружен', {
            'athletes': result,
            'total': total,
            'page': page,
            'per_page': per_page
        }

    except OperationalError as e:
        return False, f"Ошибка подключения к БД: {e}", None
    except Exception as e:
        return False, f"Неожиданная ошибка: {e}", None
    
def add_athlete_by_email(specialist_id, athlete_email):
    if not athlete_email or not athlete_email.strip():
        return False, 'Введите email спортсмена', None

    clean_email = athlete_email.strip()
    if not re.match(r'^[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+$', clean_email):
        return False, 'Неверный формат email', None

    try:
        if db.is_closed():
            db.connect()

        athlete = User.get(User.email == clean_email)
        if athlete.role != 'спортсмен':
            return False, 'Пользователь не является спортсменом', None

        specialist = User.get_by_id(specialist_id)
        specialist_role = specialist.role

        already_bound = SpecialistBinding.select().where(
            (SpecialistBinding.athlete == athlete.id) &
            (SpecialistBinding.specialist == specialist_id) &
            (SpecialistBinding.status == 'активна') &
            (SpecialistBinding.is_deleted == False)
        ).exists()
        if already_bound:
            return False, 'Этот спортсмен уже закреплён за вами', None

        bound_ids = [b.specialist_id for b in SpecialistBinding.select(SpecialistBinding.specialist_id).where(
            (SpecialistBinding.athlete == athlete.id) &
            (SpecialistBinding.status == 'активна') &
            (SpecialistBinding.is_deleted == False)
        )]

        if bound_ids:
            existing_roles = set(u.role for u in User.select(User.role).where(User.id << bound_ids))

            if specialist_role == 'тренер' and 'тренер' in existing_roles:
                return False, 'У спортсмена уже есть активный тренер', None
            if specialist_role == 'врач' and 'врач' in existing_roles:
                return False, 'У спортсмена уже есть активный врач', None

        with db.atomic():
            SpecialistBinding.create(
                athlete=athlete.id,
                specialist=specialist_id,
                status='активна',
                is_deleted=False
            )

        return True, 'Спортсмен успешно добавлен', {
            'id': athlete.id,
            'full_name': f"{athlete.last_name} {athlete.first_name}".strip(),
            'email': athlete.email
        }

    except DoesNotExist:
        return False, 'Пользователь с таким email не найден', None
    except IntegrityError:
        return False, 'Ошибка целостности данных', None
    except OperationalError as e:
        return False, f"Ошибка подключения к БД: {e}", None
def remove_athlete_from_list(specialist_id, athlete_id):
    try:
        if db.is_closed():
            db.connect()

        binding = SpecialistBinding.get(
            (SpecialistBinding.athlete == athlete_id) &
            (SpecialistBinding.specialist == specialist_id) &
            (SpecialistBinding.is_deleted == False)
        )

        with db.atomic():
            binding.status = 'прекращена'
            binding.is_deleted = True
            binding.save()

        return True, 'Спортсмен исключён из списка', None
    except DoesNotExist:
        return False, 'Привязка не найдена', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    
def get_session_recommendations(session_id):
    try:
        recommendations = Recommendation.select().where(
            (Recommendation.linked_entity == 'тренировочный план') &
            (Recommendation.linked_entity_id == session_id)
        ).order_by(Recommendation.id.desc())
        
        result = []
        for rec in recommendations:
            author = User.get_by_id(rec.author_id)
            result.append({
                'text': rec.text,
                'author_fio': f"{author.last_name} {author.first_name}".strip(),
                'author_role': author.role
            })
        
        return True, 'Рекомендации загружены', result
    except DoesNotExist:
        return False, 'Занятие не найдено', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None
    
def get_recommendations_for_entry(entry_id):
    try:
        recs = Recommendation.select().where(
            (Recommendation.linked_entity == 'дневник нагрузок') &
            (Recommendation.linked_entity_id == entry_id)
        ).order_by(Recommendation.id.desc())
        
        result = []
        for rec in recs:
            author = User.get_by_id(rec.author_id)
            result.append({
                'text': rec.text,
                'author_fio': f"{author.last_name} {author.first_name}".strip(),
                'author_role': author.role
            })
        return True, 'Загружены', result
    except Exception as e:
        return False, str(e), None

def get_diary_entries(athlete_id, start_date=None, end_date=None, page=1, per_page=3, activity_type=None):
    try:
        
        query = TrainingDiary.select().where(
            (TrainingDiary.athlete_id == athlete_id) &
            (TrainingDiary.is_deleted == False)
        )

        if start_date is not None:
            query = query.where(TrainingDiary.date >= start_date)
        if end_date is not None:
            query = query.where(TrainingDiary.date <= end_date)
        if activity_type is not None and activity_type.strip() != "":
            query = query.where(TrainingDiary.activity_type == activity_type)

        query = query.order_by(TrainingDiary.date.desc())

        total = query.count()
        offset = (page - 1) * per_page
        entries = list(query.limit(per_page).offset(offset))

        return True, "OK", {"entries": entries, "total": total}
    except Exception as e:
        print(f"Ошибка get_diary_entries: {e}")
        return False, str(e), {"entries": [], "total": 0}

def get_medical_filter_options(athlete_id):
    try:
        if db.is_closed():
            db.connect()
        types = (MedicalExam
                 .select(MedicalExam.exam_type)
                 .where(MedicalExam.athlete_id == athlete_id)
                 .distinct()
                 .order_by(MedicalExam.exam_type.asc()))
        return True, 'Ok', [t.exam_type for t in types if t.exam_type]
    except Exception as e:
        return False, str(e), []

def get_diary_filter_options(athlete_id):
    try:
        if db.is_closed():
            db.connect()
        types = (TrainingDiary
                 .select(TrainingDiary.activity_type)
                 .where((TrainingDiary.athlete_id == athlete_id) & (TrainingDiary.is_deleted == False))
                 .distinct()
                 .order_by(TrainingDiary.activity_type.asc()))
        return True, 'Ok', [t.activity_type for t in types if t.activity_type]
    except Exception as e:
        return False, str(e), []

def get_athlete_filter_options(specialist_id):
    try:
        if db.is_closed():
            db.connect()
        
        athlete_ids = list(
            SpecialistBinding.select(SpecialistBinding.athlete_id)
            .where(
                (SpecialistBinding.specialist_id == specialist_id) & 
                (SpecialistBinding.is_deleted == False)
            )
            .tuples()
        )
        if not athlete_ids:
            return True, 'Ok', {'specializations': [], 'statuses': []}
        
        athlete_ids = [id[0] for id in athlete_ids]

        specs = (User.select(User.specialization)
                 .where((User.id << athlete_ids) & (User.specialization.is_null(False)))
                 .distinct()
                 .order_by(User.specialization.asc()))
        spec_list = [s.specialization for s in specs if s.specialization]

        latest_status_subquery = (
            ReadinessStatus.select(
                ReadinessStatus.athlete_id,
                fn.MAX(ReadinessStatus.id).alias('max_id')
            )
            .where(ReadinessStatus.athlete_id << athlete_ids)
            .group_by(ReadinessStatus.athlete_id)
        )

        statuses = (ReadinessStatus.select(ReadinessStatus.current_status)
                    .where(ReadinessStatus.id << [s.max_id for s in latest_status_subquery])
                    .distinct()
                    .order_by(ReadinessStatus.current_status.asc()))
        stat_list = [s.current_status for s in statuses if s.current_status]

        return True, 'Ok', {'specializations': spec_list, 'statuses': stat_list}
    except Exception as e:
        return False, str(e), {'specializations': [], 'statuses': []}

def add_diary_recommendation(specialist_id, entry_id, text):
    clean_text = text.strip()
    if not clean_text:
        return False, 'Текст не может быть пустым', None
    if len(clean_text) > 500:
        return False, 'Текст не должен превышать 500 символов', None
    try:
        if db.is_closed():
            db.connect()

        entry = TrainingDiary.get_by_id(entry_id)
        
        if not SpecialistBinding.select().where(
            (SpecialistBinding.athlete == entry.athlete_id) &
            (SpecialistBinding.specialist == specialist_id) &
            (SpecialistBinding.status == 'активна')
        ).exists():
            return False, 'Спортсмен не закреплён за вами', None

        with db.atomic():
            Recommendation.create(
                author=specialist_id,
                athlete=entry.athlete_id,
                linked_entity='дневник нагрузок', 
                linked_entity_id=entry_id,
                text=clean_text
            )
            
        return True, 'Рекомендация сохранена', None
    except DoesNotExist:
        return False, 'Запись дневника не найдена', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def add_medical_recommendation(doctor_id, exam_id, text):
    clean_text = text.strip()
    if not clean_text:
        return False, 'Текст рекомендации не может быть пустым', None
    if len(clean_text) > 500:
        return False, 'Текст не должен превышать 500 символов', None

    try:
        if db.is_closed():
            db.connect()

        exam = MedicalExam.get_by_id(exam_id)
        
        if exam.doctor_id != doctor_id:
            return False, 'Доступ запрещён', None

        with db.atomic():
            Recommendation.create(
                author=doctor_id,
                athlete=exam.athlete_id,
                linked_entity='медкарта',
                linked_entity_id=exam_id,
                text=clean_text
            )

        return True, 'Рекомендация подписана', None
    except DoesNotExist:
        return False, 'Осмотр не найден', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def add_recommendation_to_session(specialist_id, session_id, text):
    clean_text = text.strip()
    if not clean_text:
        return False, 'Текст рекомендации не может быть пустым', None
    if len(clean_text) > 500:
        return False, 'Текст не должен превышать 500 символов', None
    try:
        if db.is_closed():
            db.connect()

        session = Session.get_by_id(session_id)
        if session.plan.coach_id != specialist_id:
            return False, 'Доступ запрещён', None
        if session.status != 'выполнено':
            return False, 'Рекомендацию можно оставить только к выполненному занятию', None

        existing = Recommendation.select().where(
            (Recommendation.author == specialist_id) &
            (Recommendation.athlete == session.plan.athlete_id) &
            (Recommendation.linked_entity == 'тренировочный план') &
            (Recommendation.linked_entity_id == session_id)
        ).first()

        if existing:
            return False, 'Рекомендация уже существует.', None

        with db.atomic():
            Recommendation.create(
                author=specialist_id,
                athlete=session.plan.athlete_id,
                linked_entity='тренировочный план', 
                linked_entity_id=session_id,
                text=clean_text
            )
        return True, 'Рекомендация сохранена', None
    except DoesNotExist:
        return False, 'Занятие не найдено', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def add_session(specialist_id, plan_id, session_date, session_time, activity_type, duration):
    if duration <= 0:
        return False, 'Длительность должна быть больше 0', None
    if session_date < date.today():
        return False, 'Дата не может быть в прошлом', None

    try:
        if db.is_closed():
            db.connect()

        plan = TrainingPlan.get_by_id(plan_id)
        if plan.coach_id != specialist_id:
            return False, 'Доступ запрещён', None

        conflict = Session.select().where(
            (Session.plan == plan_id) &
            (Session.date == session_date) &
            (Session.time == session_time) &
            (Session.is_deleted == False)
        ).exists()

        if conflict:
            return False, 'На это время уже запланировано другое занятие', None

        with db.atomic():
            Session.create(
                plan=plan_id,
                date=session_date,
                time=session_time,
                activity_type=activity_type,
                duration=duration,
                status='запланировано',
                is_deleted=False
            )
            
        return True, 'Занятие создано', None
    except DoesNotExist:
        return False, 'План не найден', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def create_medical_exam(athlete_id, doctor_id, date, exam_type, metrics):
    if date > date.today():
        return False, 'Дата осмотра не может быть в будущем', None

    try:
        if db.is_closed():
            db.connect()

        with db.atomic():
            exam = MedicalExam.create(
                athlete=athlete_id,
                doctor=doctor_id,
                exam_date=date,
                exam_type=exam_type,
                is_deleted=False
            )

            for m in metrics:
                is_crit = False
                ref = m.get('ref_range', '')
                val = m.get('value', 0)
                
                if '-' in str(ref):
                    try:
                        low, high = map(float, ref.split('-'))
                        if val < low or val > high:
                            is_crit = True
                    except ValueError:
                        pass
                
                MedicalMetric.create(
                    exam=exam,
                    metric_type=m['metric_type'],
                    value=val,
                    unit=m.get('unit', ''),
                    ref_range=ref,
                    is_critical=is_crit
                )

        return True, 'Осмотр сохранён', {'exam_id': exam.id}
    except Exception as e:
        return False, f'Ошибка: {e}', None


def create_training_plan(specialist_id, athlete_id, start_date, end_date, title='Новый план'):
    try:
        if db.is_closed():
            db.connect()

        if not SpecialistBinding.select().where(
            (SpecialistBinding.athlete == athlete_id) &
            (SpecialistBinding.specialist == specialist_id) &
            (SpecialistBinding.status == 'активна')
        ).exists():
            return False, 'Нет прав на создание плана для этого спортсмена', None

        if start_date < date.today():
            return False, 'Дата начала не может быть в прошлом', None
        if end_date <= start_date:
            return False, 'Дата окончания должна быть позже даты начала', None

        with db.atomic():
            plan = TrainingPlan.create(
                athlete=athlete_id,
                coach=specialist_id,
                title=title,
                start_date=start_date,
                end_date=end_date,
                is_deleted=False
            )
        return True, 'План создан', {'plan_id': plan.id}
    except IntegrityError as e:
        return False, 'Ошибка целостности данных', None
    except OperationalError as e:
        return False, f"Ошибка подключения: {e}", None


def delete_session(specialist_id, session_id):
    try:
        if db.is_closed():
            db.connect()

        session = Session.get_by_id(session_id)
        
        if session.plan.coach_id != specialist_id:
            return False, 'Доступ запрещён', None

        if session.date < date.today():
            return False, 'Удаление возможно только до начала даты занятия', None

        if session.status == 'выполнено':
            return False, 'Нельзя удалить выполненное занятие', None

        with db.atomic():
            session.is_deleted = True
            session.save()
        return True, 'Занятие удалено', None
    except DoesNotExist:
        return False, 'Занятие не найдено', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def delete_training_plan(specialist_id, plan_id):
    try:
        if db.is_closed():
            db.connect()
        plan = TrainingPlan.get_by_id(plan_id)
        
        if plan.coach_id != specialist_id:
            return False, 'Доступ запрещён', None

        if plan.end_date <= date.today():
            return False, 'Удаление доступно только до окончания периода плана', None

        with db.atomic():
            plan.is_deleted = True
            plan.save()
            
            Session.update(is_deleted=True).where(
                (Session.plan == plan_id) &
                (Session.is_deleted == False)
            ).execute()
            
        return True, 'План удалён', None
    except DoesNotExist:
        return False, 'План не найден', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def edit_session(specialist_id, session_id, date, time, activity_type, duration):
    if duration <= 0:
        return False, 'Длительность должна быть больше 0', None
    try:
        if db.is_closed():
            db.connect()

        session = Session.get_by_id(session_id)
        
        if session.plan.coach_id != specialist_id:
            return False, 'Доступ запрещён', None

        if session.status in ['выполнено', 'пропущено']:
            return False, 'Занятие выполнено. Редактирование невозможно', None

        plan = TrainingPlan.get_by_id(session.plan_id)
        if date < plan.start_date or date > plan.end_date:
            return False, f'Дата должна быть в рамках плана ({plan.start_date} — {plan.end_date})', None

        with db.atomic():
            session.date = date
            session.time = time
            session.activity_type = activity_type
            session.duration = duration
            session.save()
            
        return True, 'Занятие изменено', None
    except DoesNotExist:
        return False, 'Занятие не найдено', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def generate_report(specialist_id, athlete_id, report_type, start_date, end_date, fmt='excel', save_dir=None, report_name='Отчёт'):
    if save_dir is None:
        save_dir = _get_reports_dir()
    try:
        if db.is_closed():
            db.connect()

        if athlete_id is not None:
            if not SpecialistBinding.select().where(
                (SpecialistBinding.athlete == athlete_id) &
                (SpecialistBinding.specialist == specialist_id) &
                (SpecialistBinding.status == 'активна')
            ).exists():
                return False, 'Спортсмен не закреплён за вами', None
            athlete_ids = [athlete_id]
        else:
            bindings = SpecialistBinding.select().where(
                (SpecialistBinding.specialist == specialist_id) &
                (SpecialistBinding.status == 'активна')
            )
            athlete_ids = [b.athlete_id for b in bindings]
            if not athlete_ids:
                return False, 'У вас нет закреплённых спортсменов', None

        os.makedirs(save_dir, exist_ok=True)
        
        suffix = athlete_id if athlete_id else 'all'
        ext = 'xlsx' if fmt == 'excel' else 'pdf'
        safe_name = report_name.replace('/', '_').replace('\\', '_').replace(':', '_')
        filename = f"{save_dir}/{safe_name}_{suffix}_{date.today().strftime('%Y%m%d')}.{ext}"

        include_training = report_type in ('general', 'training')
        include_medical  = report_type in ('general', 'medical')
        include_diary    = report_type in ('general', 'diary')

        def get_athlete_name(aid):
            try:
                u = User.get_by_id(aid)
                return f"{u.last_name} {u.first_name}"
            except Exception:
                return str(aid)

        training_rows = []
        medical_rows  = []
        diary_rows    = []
        period_str = f"Период: {start_date} — {end_date}"

        for aid in athlete_ids:
            aname = get_athlete_name(aid)

            if include_training:
                plans = TrainingPlan.select().where(
                    (TrainingPlan.athlete == aid) &
                    (TrainingPlan.start_date <= end_date) &
                    (TrainingPlan.end_date >= start_date)
                )
                for plan in plans:
                    sessions = Session.select().where(
                        (Session.plan == plan) & (Session.is_deleted == False) &
                        (Session.date >= start_date) & (Session.date <= end_date)
                    )
                    for s in sessions:
                        training_rows.append([
                            aname, str(s.date), s.activity_type or '',
                            s.duration, s.status or ''
                        ])

            if include_medical:
                exams = MedicalExam.select().where(
                    (MedicalExam.athlete == aid) &
                    (MedicalExam.exam_date >= start_date) &
                    (MedicalExam.exam_date <= end_date)
                )
                for exam in exams:
                    metrics = MedicalMetric.select().where(MedicalMetric.exam == exam)
                    for m in metrics:
                        critical_text = 'Да' if m.is_critical else 'Нет'
                        medical_rows.append([
                            aname, str(exam.exam_date), exam.exam_type or '',
                            m.metric_type or '',
                            str(m.value) if m.value is not None else '',
                            m.unit or '',
                            critical_text 
                        ])

            if include_diary:
                entries = TrainingDiary.select().where(
                    (TrainingDiary.athlete == aid) &
                    (TrainingDiary.date >= start_date) &
                    (TrainingDiary.date <= end_date) &
                    (TrainingDiary.is_deleted == False)
                )
                for d in entries:
                    diary_rows.append([
                        aname, str(d.date), d.activity_type or '', d.duration,
                        d.steps, d.sleep_hours, d.fatigue, d.mood
                    ])

        if not training_rows and not medical_rows and not diary_rows:
            return False, 'Данных за выбранный период не найдено.', None

        if fmt == 'excel':
            try:
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            except ImportError:
                return False, 'Установите openpyxl: pip install openpyxl', None

            wb = Workbook()
            wb.remove(wb.active)
            header_font_white = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1a1a1a")
            title_font = Font(bold=True, size=14, color="1a1a1a")
            period_font = Font(size=11, color="555555")
            critical_fill = PatternFill("solid", fgColor="ffcccc")
            critical_font = Font(color="cc0000", bold=True)

            def make_sheet(ws, headers, rows, title_text, period_text, critical_col_idx=None):
                from openpyxl.utils import get_column_letter
                last_col = len(headers)
                
                ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
                title_cell = ws.cell(row=1, column=1, value=title_text)
                title_cell.font = title_font
                title_cell.alignment = Alignment(horizontal='center')

                ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
                period_cell = ws.cell(row=2, column=1, value=period_text)
                period_cell.font = period_font
                period_cell.alignment = Alignment(horizontal='center')

                for col_idx, h in enumerate(headers, 1):
                    cell = ws.cell(row=3, column=col_idx, value=h)
                    cell.font = header_font_white
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = Border(bottom=Side(style='medium'))

                for r_idx, row_data in enumerate(rows, 4):
                    for c_idx, val in enumerate(row_data, 1):
                        cell = ws.cell(row=r_idx, column=c_idx, value=val)
                        cell.alignment = Alignment(vertical='top', wrap_text=True)

                        if critical_col_idx is not None:
                            if c_idx == critical_col_idx and val == 'Да':
                                for highlight_c in range(1, last_col + 1):
                                    highlight_cell = ws.cell(row=r_idx, column=highlight_c)
                                    highlight_cell.fill = critical_fill
                                    highlight_cell.font = critical_font

                for c in range(1, last_col + 1):
                    col_letter = get_column_letter(c)
                    max_len = 0
                    for row in ws.iter_rows(min_col=c, max_col=c, min_row=3):
                        for cell in row:
                            if cell.value:
                                max_len = max(max_len, len(str(cell.value)))
                    width = min(max(max_len * 1.3 + 2, 12), 45)
                    ws.column_dimensions[col_letter].width = width

                ws.freeze_panes = 'A4'

            if include_training and training_rows:
                ws = wb.create_sheet('Тренировки')
                make_sheet(ws, ['Спортсмен', 'Дата', 'Активность', 'Длительность (мин)', 'Статус'], 
                          training_rows, report_name, period_str)

            if include_medical and medical_rows:
                ws = wb.create_sheet('Медосмотры')
                make_sheet(ws, ['Спортсмен', 'Дата', 'Тип осмотра', 'Показатель', 'Значение', 'Ед.изм.', 'Критично'], 
                          medical_rows, report_name, period_str, critical_col_idx=7)

            if include_diary and diary_rows:
                ws = wb.create_sheet('Дневник нагрузок')
                make_sheet(ws, ['Спортсмен', 'Дата', 'Активность', 'Длительность', 'Шаги', 'Сон (ч)', 'Усталость', 'Настроение'], 
                          diary_rows, report_name, period_str)

            if not wb.worksheets:
                ws = wb.create_sheet('Нет данных')
                ws.cell(row=1, column=1, value='Нет данных за выбранный период')

            wb.save(filename)

        elif fmt == 'pdf':
            try:
                from reportlab.lib.pagesizes import A4, landscape
                from reportlab.platypus import SimpleDocTemplate, Table, Paragraph, Spacer, PageBreak
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.lib import colors
                from reportlab.pdfbase import pdfmetrics
                from reportlab.pdfbase.ttfonts import TTFont
            except ImportError:
                return False, 'Установите reportlab: pip install reportlab', None

            font_registered = False
            font_name = 'Helvetica'
            candidates = [
                ('DejaVu', '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
                ('DejaVu', '/usr/share/fonts/dejavu/DejaVuSans.ttf'),
                ('Arial',  'C:/Windows/Fonts/arial.ttf'),
                ('Arial',  '/Library/Fonts/Arial.ttf'),
                ('FreeSans', '/usr/share/fonts/gnu-free/FreeSans.ttf'),
            ]
            for fname, fpath in candidates:
                if os.path.exists(fpath):
                    try:
                        pdfmetrics.registerFont(TTFont(fname, fpath))
                        font_name = fname
                        font_registered = True
                        break
                    except Exception:
                        continue

            if not font_registered:
                try:
                    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
                    pdfmetrics.registerFont(UnicodeCIDFont('HeiseiMin-W3'))
                    font_name = 'HeiseiMin-W3'
                    font_registered = True
                except Exception:
                    pass

            page = landscape(A4) if report_type == 'general' else A4
            doc = SimpleDocTemplate(filename, pagesize=page, leftMargin=36, rightMargin=36, topMargin=40, bottomMargin=36)
            W = doc.width

            h1_style = ParagraphStyle('H1', fontName=font_name, fontSize=18, leading=22, spaceAfter=4, textColor=colors.HexColor('#1a1a1a'), alignment=1)
            h2_style = ParagraphStyle('H2', fontName=font_name, fontSize=11, leading=17, spaceAfter=1, textColor=colors.HexColor('#555555'))
            h3_style = ParagraphStyle('H3', fontName=font_name, fontSize=14, leading=18, spaceAfter=8, textColor=colors.HexColor('#1a1a1a'), fontWeight='bold')

            def make_table(headers, rows, col_ratios=None, critical_col_idx=None):
                col_w = [W * r for r in col_ratios] if col_ratios else [W / len(headers)] * len(headers)
                data = [headers] + rows if rows else [headers, ['Нет данных']]
                t = Table(data, colWidths=col_w, repeatRows=1)
                
                style = [
                    ('FONTNAME', (0, 0), (-1, -1), font_name),
                    ('FONTSIZE', (0, 0), (-1, 0), 10),
                    ('FONTSIZE', (0, 1), (-1, -1), 9),
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1a1a')),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                    ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                    ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
                    ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#cccccc')),
                    ('LEFTPADDING', (0, 0), (-1, -1), 6),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                    ('WORDWRAP', (0, 0), (-1, -1), 'CJK')
                ]
                
                if critical_col_idx is not None:
                    for row_idx, row_data in enumerate(rows, 1): 
                        if len(row_data) >= critical_col_idx and row_data[critical_col_idx - 1] == 'Да':
                            style.append(('BACKGROUND', (0, row_idx), (-1, row_idx), colors.HexColor('#ffcccc')))
                            style.append(('TEXTCOLOR', (0, row_idx), (-1, row_idx), colors.HexColor('#cc0000')))
                
                t.setStyle(style)
                return t

            elements = []
            elements.append(Paragraph(report_name, h1_style))
            elements.append(Paragraph(period_str, h2_style))
            elements.append(Spacer(1, 12))

            if include_training and training_rows:
                elements.append(Paragraph('Тренировки', h3_style))
                t_rows = [[r[0], r[1], r[2], str(r[3]), r[4]] for r in training_rows]
                elements.append(make_table(['Спортсмен', 'Дата', 'Активность', 'Длительность (мин)', 'Статус'], 
                                         t_rows, col_ratios=[0.20, 0.12, 0.32, 0.18, 0.18]))
                if report_type == 'general':
                    elements.append(PageBreak())
                else:
                    elements.append(Spacer(1, 16))

            if include_medical and medical_rows:
                elements.append(Paragraph('Медицинские показатели', h3_style))
                m_rows = [[r[0], r[1], r[2], r[3], r[4], r[5], r[6]] for r in medical_rows]
                elements.append(make_table(['Спортсмен', 'Дата', 'Тип осмотра', 'Показатель', 'Значение', 'Ед.изм.', 'Критично'], 
                                         m_rows, col_ratios=[0.16, 0.10, 0.18, 0.22, 0.10, 0.10, 0.14], 
                                         critical_col_idx=7))
                if report_type == 'general' and include_diary:
                    elements.append(PageBreak())
                elif report_type != 'general':
                    elements.append(Spacer(1, 16))

            if include_diary and diary_rows:
                elements.append(Paragraph('Дневник нагрузок', h3_style))
                d_rows = [[r[0], r[1], r[2], str(r[3]), str(r[4]), str(r[5]), str(r[6]), str(r[7])] for r in diary_rows]
                elements.append(make_table(['Спортсмен', 'Дата', 'Активность', 'Длительность', 'Шаги', 'Сон', 'Усталость', 'Настроение'], 
                                         d_rows, col_ratios=[0.14, 0.10, 0.20, 0.11, 0.11, 0.08, 0.08, 0.18]))

            try:
                doc.build(elements)
            except Exception as pdf_err:
                return False, f'Ошибка генерации PDF: {pdf_err}', None
        else:
            return False, 'Неподдерживаемый формат. Используйте excel или pdf', None

        return True, f'Отчёт сохранён: {filename}', {'path': filename}

    except OperationalError as e:
        return False, f'Ошибка БД: {e}', None
    except Exception as e:
        import traceback; traceback.print_exc()
        return False, f'Ошибка генерации: {e}', None


def get_athlete_medical_records(doctor_id, athlete_id, exam_date=None, exam_type=None):
    try:
        if db.is_closed():
            db.connect()

        if not SpecialistBinding.select().where(
            (SpecialistBinding.athlete == athlete_id) &
            (SpecialistBinding.specialist == doctor_id) &
            (SpecialistBinding.status == 'активна')
        ).exists():
            return False, 'Спортсмен не закреплён за вами', None

        conditions = [MedicalExam.athlete == athlete_id]

        if exam_date:
            conditions.append(MedicalExam.exam_date == exam_date)

        if exam_type:
            conditions.append(MedicalExam.exam_type == exam_type)

        exams = MedicalExam.select().where(*conditions).order_by(MedicalExam.exam_date.desc())

        result = []
        for exam in exams:
            metrics = MedicalMetric.select().where(
                MedicalMetric.exam == exam
            ).order_by(MedicalMetric.id.desc())

            recommendations = Recommendation.select().where(
                (Recommendation.linked_entity == 'медкарта') &
                (Recommendation.linked_entity_id == exam.id)
            ).order_by(Recommendation.id.desc())

            try:
                doctor = exam.doctor
                doctor_fio = f"{doctor.last_name} {doctor.first_name} {doctor.middle_name or ''}".strip()
                doctor_email = doctor.email
            except DoesNotExist:
                doctor_fio = "Врач удалён"
                doctor_email = "—"

            result.append({
                'exam_id': exam.id,
                'exam_date': exam.exam_date,
                'exam_type': exam.exam_type,
                'doctor_fio': doctor_fio,
                'doctor_email': doctor_email,
                'metrics': [{
                    'id': m.id,
                    'type': m.metric_type,
                    'value': m.value,
                    'unit': m.unit,
                    'ref_range': m.ref_range or '',
                    'is_critical': bool(m.is_critical)
                } for m in metrics],
                'recommendations': [{
                    'id': r.id,
                    'text': r.text
                } for r in recommendations]
            })

        return True, 'Медицинские данные загружены', result
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


def update_athlete_status(specialist_id, athlete_id, new_status):
    clean_status = unicodedata.normalize('NFKC', new_status.strip().lower())
    valid_statuses = ['здоров', 'устал', 'болен']
    
    if clean_status not in valid_statuses:
        return False, f'Недопустимый статус: "{clean_status}"', None

    try:
        if db.is_closed():
            db.connect()

        if not SpecialistBinding.select().where(
            (SpecialistBinding.athlete == athlete_id) &
            (SpecialistBinding.specialist == specialist_id) &
            (SpecialistBinding.status == 'активна')
        ).exists():
            return False, 'Спортсмен не закреплён за вами', None

        specialist = User.get_by_id(specialist_id)
        specialist_role = specialist.role

        try:
            last_status = ReadinessStatus.select().where(
                ReadinessStatus.athlete == athlete_id
            ).order_by(ReadinessStatus.id.desc()).get()

            if last_status.initiator.role == 'врач':
                if specialist_role == 'тренер' and last_status.current_status in ['болен', 'устал']:
                    return False, 'Статус заблокирован врачом (спортсмен болен или устал).', None
                
        except DoesNotExist:
            pass 

        with db.atomic():
            lock_status = 'заблокировано' if specialist_role == 'врач' else 'свободно'
            
            ReadinessStatus.create(
                athlete=athlete_id,
                initiator=specialist_id,
                current_status=clean_status,
                lock_status=lock_status  
            )
            
        return True, 'Статус обновлён', None
    except OperationalError as e:
        return False, f"Ошибка БД: {e}", None


